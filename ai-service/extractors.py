import pdfplumber
from docx import Document
from pptx import Presentation
import os
from pathlib import Path


PROJECT_ROOT = Path(__file__).resolve().parents[1]
DEFAULT_DOCUMENT_ROOTS = [
    PROJECT_ROOT / "storage" / "app" / "private" / "resources",
    PROJECT_ROOT / "storage" / "app" / "private" / "temp",
]


def _resolve_document_path(file_path: str) -> str:
    path = Path(file_path).expanduser().resolve(strict=True)
    if not path.is_file():
        raise FileNotFoundError("Fichier introuvable")

    configured_roots = os.getenv("AI_DOCUMENT_ROOTS")
    roots = (
        [Path(root).expanduser().resolve() for root in configured_roots.split(os.pathsep) if root]
        if configured_roots
        else DEFAULT_DOCUMENT_ROOTS
    )

    for root in roots:
        try:
            path.relative_to(root)
            return str(path)
        except ValueError:
            continue

    raise PermissionError("Chemin de document non autorisé")


def extract_text(file_path: str, file_type: str) -> str:
    """
    Extrait le texte brut d'un document selon son type.
    file_type doit être l'un de : 'pdf', 'word', 'powerpoint'
    """
    file_path = _resolve_document_path(file_path)

    if file_type == "pdf":
        return _extract_from_pdf(file_path)
    elif file_type == "word":
        return _extract_from_word(file_path)
    elif file_type == "powerpoint":
        return _extract_from_powerpoint(file_path)
    else:
        raise ValueError(f"Type de fichier non supporté : {file_type}")


def _extract_from_pdf(file_path: str) -> str:
    text_parts = []
    with pdfplumber.open(file_path) as pdf:
        for page_number, page in enumerate(pdf.pages, start=1):
            page_text = page.extract_text()
            if page_text:
                text_parts.append(f"[Page {page_number}]\n{page_text}")
    return "\n\n".join(text_parts)


def _extract_from_word(file_path: str) -> str:
    doc = Document(file_path)
    paragraphs = [p.text for p in doc.paragraphs if p.text.strip()]
    return "\n".join(paragraphs)


def _extract_from_powerpoint(file_path: str) -> str:
    prs = Presentation(file_path)
    text_parts = []
    for slide_number, slide in enumerate(prs.slides, start=1):
        slide_text = []
        for shape in slide.shapes:
            if hasattr(shape, "text") and shape.text.strip():
                slide_text.append(shape.text)
        if slide_text:
            text_parts.append(f"[Slide {slide_number}]\n" + "\n".join(slide_text))
    return "\n\n".join(text_parts)
